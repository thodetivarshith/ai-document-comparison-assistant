"""Build the downloadable DOCX report from a validated comparison result."""
from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.shared import Pt

from comparison import evidence_badge
from document_reader import DocumentData

DISCLAIMER = ("This is an informational, AI-generated comparison based only on the extracted text. "
              "It is not legal advice and does not certify that either document is compliant, safe or equivalent.")


def _table(doc, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, h in zip(table.rows[0].cells, headers):
        cell.text = ""
        cell.paragraphs[0].add_run(h).bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)


def _evidence(text: str, verified) -> str:
    if not text:
        return ""
    badge = evidence_badge(verified)
    return f'"{text}"' + (f" [{badge}]" if badge else "")


def build_report(result: dict, doc_a: DocumentData, doc_b: DocumentData) -> bytes:
    doc = Document()
    doc.add_heading("AI Document Comparison Report", 0)
    doc.add_paragraph(f"Generated: {datetime.now():%Y-%m-%d %H:%M}")
    p = doc.add_paragraph(DISCLAIMER)
    p.runs[0].italic = True

    for label, d, key in (("Document A", doc_a, "document_a"), ("Document B", doc_b, "document_b")):
        doc.add_heading(f"{label} — {d.name}", 1)
        doc.add_paragraph(f"{d.page_count} page(s), {d.char_count:,} characters extracted.")
        doc.add_paragraph(result[key]["summary"])
        for point in result[key]["main_points"]:
            doc.add_paragraph(point, style="List Bullet")

    doc.add_heading("Similarities", 1)
    if result["similarities"]:
        _table(doc, ["Topic", "Detail", "Evidence A", "Evidence B"], [
            [s["topic"], s["detail"],
             _evidence(s["evidence_a"], s.get("evidence_a_verified")),
             _evidence(s["evidence_b"], s.get("evidence_b_verified"))]
            for s in result["similarities"]])
    else:
        doc.add_paragraph("No similarities were identified.")

    doc.add_heading("Differences", 1)
    if result["differences"]:
        _table(doc, ["Topic", "Type", "Document A", "Document B", "Why it matters"], [
            [d["topic"], d["type"],
             d["document_a"] + (("\n" + _evidence(d["evidence_a"], d.get("evidence_a_verified"))) if d["evidence_a"] else ""),
             d["document_b"] + (("\n" + _evidence(d["evidence_b"], d.get("evidence_b_verified"))) if d["evidence_b"] else ""),
             d["significance"]]
            for d in result["differences"]])
    else:
        doc.add_paragraph("No differences were identified.")

    doc.add_heading("Key Details Comparison", 1)
    if result["key_details"]:
        _table(doc, ["Category", "Item", "Document A", "Document B", "Status"], [
            [k["category"], k["item"], k["document_a"], k["document_b"], k["status"].replace("_", " ")]
            for k in result["key_details"]])
    else:
        doc.add_paragraph("No key details were identified.")

    doc.add_heading("Overall Conclusion", 1)
    doc.add_paragraph(result["conclusion"])

    doc.add_heading("Limitations", 1)
    limits = result["limitations"] or ["No specific limitations were reported."]
    for item in limits:
        doc.add_paragraph(item, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_markdown(result: dict, doc_a: DocumentData, doc_b: DocumentData) -> str:
    """Plain-text/Markdown export with the same sections as the DOCX report."""
    def esc(s: str) -> str:
        return (s or "").replace("|", "\\|").replace("\n", " ")

    out = ["# AI Document Comparison Report", f"_Generated {datetime.now():%Y-%m-%d %H:%M}_", "", f"> {DISCLAIMER}", ""]
    for label, d, key in (("Document A", doc_a, "document_a"), ("Document B", doc_b, "document_b")):
        out += [f"## {label} — {d.name}", f"{d.page_count} page(s), {d.char_count:,} characters extracted.", "",
                result[key]["summary"], ""]
        out += [f"- {p}" for p in result[key]["main_points"]] + [""]
    out += ["## Similarities", ""]
    out += [f"- **{s['topic']}** — {s['detail']}" for s in result["similarities"]] or ["No similarities identified."]
    out += ["", "## Differences", ""]
    for d in result["differences"]:
        out += [f"### {d['topic']} ({d['type']})", f"- **A:** {d['document_a']}", f"- **B:** {d['document_b']}"]
        if d["significance"]:
            out.append(f"- _Why it matters:_ {d['significance']}")
        out.append("")
    if not result["differences"]:
        out.append("No differences identified.")
    out += ["", "## Key Details Comparison", "", "| Category | Item | Document A | Document B | Status |", "|---|---|---|---|---|"]
    out += [f"| {esc(k['category'])} | {esc(k['item'])} | {esc(k['document_a'])} | {esc(k['document_b'])} | "
            f"{k['status'].replace('_', ' ')} |" for k in result["key_details"]]
    out += ["", "## Overall Conclusion", "", result["conclusion"], "", "## Limitations", ""]
    out += [f"- {x}" for x in result["limitations"]] or ["- No specific limitations were reported."]
    return "\n".join(out) + "\n"
